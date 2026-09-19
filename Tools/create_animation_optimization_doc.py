from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"E:\Unity\AwCon\Docs\Unity动画资源瘦身操作步骤.docx"

def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        tbl_borders.append(e)
    tbl_pr.append(tbl_borders)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, 15 if level == 1 else 12, True)
    return p

def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(text)
    set_font(r, 10.5)
    return p

def add_step(doc, title, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title + "：")
    set_font(r, 10.5, True)
    r = p.add_run(text)
    set_font(r, 10.5)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.3)
sec.right_margin = Cm(2.3)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
r = title.add_run("Unity 动画资源瘦身操作步骤")
set_font(r, 22, True)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(18)
r = sub.add_run("FBX 动画迁移为独立 ANIM 并安全清理演示资源")
set_font(r, 10.5, False, (80, 80, 80))

add_body(doc, "本文记录本项目已执行的动画资源优化流程。目标是在不丢失玩家、AI 与动画控制器引用的前提下，删除未使用资源，并将仍在使用的 FBX 动画迁移为独立 .anim 文件。完成后，FemaleMovementAnimsetPro 已删除，Emotion 已从 38.24MB 缩小到约 9.25MB，FemaleRunner 已改为仅保留独立动画文件。")

add_heading(doc, "一 适用场景")
add_body(doc, "适用于 Asset Store 动画包、角色演示包或历史资源包。特别适合目录中混有 FBX、模型、贴图、场景、控制器与脚本，但项目实际只需要部分动画动作的情况。")

add_heading(doc, "二 重要原则")
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
borders(table)
headers = ["原则", "原因"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    shade(cell, "1F4E78")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(text)
    set_font(rr, 10, True, (255, 255, 255))
for a, b in [
    ("先检查依赖", "直接删除 FBX 可能让 Animator Controller、ScriptableObject、Prefab 或场景丢失动画引用。"),
    ("保持 GUID", "移动仍需要的贴图或资源时，使用 Unity 的 AssetDatabase.MoveAsset，材质和 Prefab 引用才不会丢失。"),
    ("先迁移后删除", "先创建独立 .anim 并替换引用，再验证没有外部依赖，最后删除原始目录。"),
    ("在 Unity 内执行", "不要用资源管理器直接删除 Unity 资源；应通过 Unity Project 窗口或 Editor API 删除，让 .meta 与引用一起处理。"),
]:
    row = table.add_row().cells
    for idx, text in enumerate((a, b)):
        row[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = row[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        rr = p.add_run(text)
        set_font(rr, 9.5, idx == 0)

add_heading(doc, "三 操作总流程")
add_step(doc, "确认目录大小和资源类型", "统计目录内 FBX、.anim、贴图、模型、场景与脚本的数量和体积。优先处理占用大且明显为演示内容的目录。")
add_step(doc, "查找项目外部依赖", "使用 Unity 的 AssetDatabase.GetDependencies 扫描所有场景、Prefab、Animator Controller、ScriptableObject 和材质，确认目标目录被哪些资源使用。")
add_step(doc, "判断处理方式", "没有外部依赖的动画包可整体删除。仍被使用的 FBX 动画必须先复制为独立 .anim，并替换现有引用。")
add_step(doc, "创建独立动画", "在目标文件夹中创建 AnimationClip，通过 EditorUtility.CopySerialized 复制 FBX 内的 AnimationClip。这样可保留动画曲线、循环设置与事件信息。")
add_step(doc, "替换引用", "ScriptableObject 中的 AnimationClip 字段用 SerializedObject 替换；Animator Controller 中的 AnimatorState.motion 也要指向新的 .anim。")
add_step(doc, "迁移共用资源", "若角色 Prefab 或材质仍引用旧动画包中的贴图，先将贴图移动到角色资源目录，保持原 GUID 不变。")
add_step(doc, "验证后删除", "再次扫描外部依赖，确认数量为 0 后删除旧目录。最后将新的 AnimOnly 目录改回原目录名称，保持项目结构清晰。")
add_step(doc, "重新编译验证", "等待 Unity 导入完成，检查 Compilation Errors 和 Console。发现错误时先恢复或补齐引用，不要继续提交。")

add_heading(doc, "四 本项目实际处理记录")
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
borders(table)
for i, text in enumerate(["目录", "处理前", "处理方式", "结果"]):
    c = table.rows[0].cells[i]
    shade(c, "1F4E78")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(text)
    set_font(rr, 9.5, True, (255, 255, 255))
records = [
    ("FemaleRunner", "211.49MB，216 个 FBX", "提取 94 个 .anim，替换 Player 和 AI 的 190 处引用；迁移两张共用贴图", "仅保留 94 个 .anim，约 77.58MB"),
    ("FemaleMovementAnimsetPro", "118.33MB", "依赖扫描结果为 0，整体删除", "不影响当前项目"),
    ("Emotion", "38.24MB，7 个 FBX", "复制 7 个独立 .anim，替换 Shotgun_Controller 的 7 个状态", "约 9.25MB，控制器引用正常"),
]
for rec in records:
    cells = table.add_row().cells
    for i, text in enumerate(rec):
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        rr = p.add_run(text)
        set_font(rr, 9)

add_heading(doc, "五 Unity Editor API 示例")
add_body(doc, "以下示例展示核心思路。实际执行时应加入路径存在性检查、错误处理、保存与读回验证。")
code = '''// 1. 将 FBX 中的 AnimationClip 复制为独立 .anim\nvar source = AssetDatabase.LoadAssetAtPath<AnimationClip>(fbxPath);\nvar copy = new AnimationClip();\nEditorUtility.CopySerialized(source, copy);\nAssetDatabase.CreateAsset(copy, animPath);\n\n// 2. 替换 Animator Controller 状态引用\nstate.motion = AssetDatabase.LoadAssetAtPath<AnimationClip>(animPath);\nEditorUtility.SetDirty(controller);\n\n// 3. 保持 GUID 地移动贴图\nstring error = AssetDatabase.MoveAsset(oldTexturePath, newTexturePath);\n\n// 4. 最后删除已无依赖的原始目录\nbool deleted = AssetDatabase.DeleteAsset(oldFolderPath);\nAssetDatabase.SaveAssets();\nAssetDatabase.Refresh();'''
p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(0.45)
p.paragraph_format.space_after = Pt(6)
run = p.add_run(code)
set_font(run, 8.5)
run.font.name = "Consolas"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")

add_heading(doc, "六 验证清单")
for item in [
    "目标目录外部依赖扫描结果为 0。",
    "Animator Controller、ScriptableObject、Prefab 与场景没有 Missing 引用。",
    "新 .anim 数量与实际需要的动作数量一致。",
    "Unity Compilation Errors 为 0。",
    "Console 没有新增错误。",
    "在 Play Mode 测试走、跑、冲刺、起步、停止、跳跃、翻滚和表情状态。",
    "确认 Git diff 只包含计划内的资源删除、.anim 新增和控制器/配置引用更新。",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(item)
    set_font(r, 10)

add_heading(doc, "七 注意事项")
add_body(doc, "独立 .anim 本身仍可能较大，因为它保存完整骨骼关键帧。若要继续压缩，应先保留所需动作，再对关键帧做保守误差压缩，并在 Play Mode 中观察腿部、手部与根运动是否出现抖动或滑步。ZIP 仅适合交付或备份，不会降低 Unity 项目内资源的运行时占用。")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Unity 动画资源瘦身操作步骤")
set_font(r, 8, False, (110, 110, 110))

doc.save(OUT)
print(OUT)
